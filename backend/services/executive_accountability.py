"""Constant-query executive accountability read model for TASK_209."""

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import text

from api.dependencies import normalize_role
from services import observation_quality
from services.field_attention import MAX_SCOPE_FIELDS, build_attention_queue


TASHKENT = ZoneInfo("Asia/Tashkent")

# "Latest accepted stored observation" uses the canonical quality contract in
# services/observation_quality.py, so a NULL cloud_cover_pct written by the
# canonical collectors no longer hides an otherwise good observation.
_ACCEPTED_NDVI_OBSERVATION = observation_quality.accepted_observation_sql(
    value_column="n.mean_ndvi",
    valid_pixels_column="n.valid_pixels_pct",
    cloud_column="n.cloud_cover_pct",
    minimum_valid_pixels_pct=observation_quality.MIN_VALID_PIXELS_ANALYSIS_PCT,
)
_ACCEPTED_INDEX_OBSERVATION = observation_quality.accepted_observation_sql(
    value_column="s.mean_value",
    valid_pixels_column="s.valid_pixels_pct",
    cloud_column="s.cloud_cover_pct",
    minimum_valid_pixels_pct=observation_quality.MIN_VALID_PIXELS_ANALYSIS_PCT,
)
DEFINITIONS_VERSION = "task225_canonical_backlog_v1"

# Current-status definitions (TASK_225). Open inspections use the canonical
# TASK_217 open states plus still-open legacy rows; open actions are the
# current-cycle work of a live TASK_220 plan; awaiting verification is a plan in
# pending_verification. The retired corrective_actions lifecycle no longer
# feeds any current-status count.
OPEN_INSPECTION_STATUSES = "('pending','new','assigned','in_progress','submitted')"
LIVE_PLAN = "p.status NOT IN ('closed','cancelled','superseded')"
ACTIVE_WORK = "w.status IN ('planned','in_progress')"
WORK_DUE_DATE = "(w.due_at AT TIME ZONE 'Asia/Tashkent')::date"
MANAGEMENT_ROLES = frozenset({"admin", "manager"})
ACCOUNTABILITY_KINDS = frozenset(
    {
        "unassigned_inspections",
        "overdue_inspections",
        "open_actions",
        "overdue_actions",
        "awaiting_verification",
    }
)
LIMITATIONS = [
    (
        "attention_signal_to_inspection_hours starts at local midnight on the "
        "source observation date because the current schema has no durable "
        "attention_entered_at event."
    ),
    (
        "Satellite verification reports observed index direction and does not "
        "prove agronomic causality."
    ),
    (
        "No-data, stale-data, and low-confidence counts are operational data "
        "quality warnings, not agronomic diagnoses."
    ),
    (
        "Backlog, owner and accountability counts read the canonical inspection "
        "and agronomy-plan lifecycles. Cycle-time and verification-outcome "
        "metrics still describe the retired TASK_209 corrective-action history "
        "until the management analytics definitions are re-baselined."
    ),
]


@dataclass(frozen=True, slots=True)
class ExecutiveScope:
    role: str
    user_id: int
    enterprise_id: int | None


@dataclass(frozen=True, slots=True)
class DateWindow:
    date_from: date
    date_to: date
    from_timestamp: datetime
    to_exclusive: datetime


def resolve_scope(user, requested_enterprise_id=None) -> ExecutiveScope:
    role = normalize_role(user)
    if role not in MANAGEMENT_ROLES:
        raise HTTPException(403, "Executive accountability requires management role")
    if role == "manager":
        if user.enterprise_id is None:
            raise HTTPException(403, "Manager has no enterprise_id")
        if (
            requested_enterprise_id is not None
            and requested_enterprise_id != user.enterprise_id
        ):
            raise HTTPException(403, "Enterprise is outside manager scope")
        enterprise_id = user.enterprise_id
    else:
        enterprise_id = requested_enterprise_id
    return ExecutiveScope(
        role=role,
        user_id=user.id,
        enterprise_id=enterprise_id,
    )


def resolve_window(date_from=None, date_to=None, *, today=None) -> DateWindow:
    resolved_to = date_to or today or datetime.now(TASHKENT).date()
    resolved_from = date_from or (resolved_to - timedelta(days=29))
    if resolved_from > resolved_to:
        raise HTTPException(422, "date_from must not be later than date_to")
    if (resolved_to - resolved_from).days + 1 > 366:
        raise HTTPException(422, "date range exceeds 366 inclusive days")
    return DateWindow(
        date_from=resolved_from,
        date_to=resolved_to,
        from_timestamp=datetime.combine(resolved_from, time.min, TASHKENT),
        to_exclusive=datetime.combine(
            resolved_to + timedelta(days=1),
            time.min,
            TASHKENT,
        ),
    )


def _one(result):
    if hasattr(result, "mappings"):
        return result.mappings().first()
    row = result.fetchone()
    return row._mapping if row is not None and hasattr(row, "_mapping") else row


def _all(result):
    if hasattr(result, "mappings"):
        return list(result.mappings().all())
    return [
        row._mapping if hasattr(row, "_mapping") else row
        for row in result.fetchall()
    ]


def _integer(value):
    return int(value or 0)


def _hours(value):
    return None if value is None else round(float(value), 2)


def _duration(row, prefix):
    return {
        "sample_count": _integer(row[f"{prefix}_count"]),
        "median_hours": _hours(row[f"{prefix}_median"]),
        "p90_hours": _hours(row[f"{prefix}_p90"]),
    }


def _overview_sql(tenant_clause):
    return text(
        f"""
        WITH scope_fields AS (
          SELECT f.id AS field_id, f.enterprise_id, e.name AS enterprise_name
          FROM fields f
          JOIN enterprises e ON e.id=f.enterprise_id
          {tenant_clause}
        ),
        inspection_snapshot AS (
          SELECT
            count(*) FILTER (
              WHERE i.status IN {OPEN_INSPECTION_STATUSES}
            ) AS open_inspections,
            count(*) FILTER (
              WHERE i.status IN {OPEN_INSPECTION_STATUSES}
                AND i.assigned_to_id IS NULL
            ) AS unassigned_inspections,
            count(*) FILTER (
              WHERE i.status IN {OPEN_INSPECTION_STATUSES}
                AND i.due_date < :as_of_date
            ) AS overdue_inspections
          FROM field_inspections i
          JOIN scope_fields sf ON sf.field_id=i.field_id
        ),
        live_work AS (
          SELECT w.id,w.field_id,w.assigned_to_id,{WORK_DUE_DATE} AS due_date
          FROM agronomy_work_items w
          JOIN agronomy_plans p ON p.id=w.plan_id AND p.enterprise_id=w.enterprise_id
            AND p.cycle=w.cycle
          JOIN scope_fields sf ON sf.field_id=w.field_id
          WHERE {LIVE_PLAN} AND {ACTIVE_WORK}
        ),
        action_snapshot AS (
          SELECT
            count(*) AS open_actions,
            count(*) FILTER (WHERE due_date < :as_of_date) AS overdue_actions
          FROM live_work
        ),
        verification_snapshot AS (
          SELECT count(*) AS awaiting_verification
          FROM agronomy_plans p
          JOIN scope_fields sf ON sf.field_id=p.field_id
          WHERE p.status='pending_verification'
        ),
        attention_durations AS (
          SELECT extract(
            epoch FROM (
              i.created_at
              - (i.source_observation_date::timestamp AT TIME ZONE 'Asia/Tashkent')
            )
          ) / 3600.0 AS hours
          FROM field_inspections i
          JOIN scope_fields sf ON sf.field_id=i.field_id
          WHERE i.source='attention_queue'
            AND i.source_observation_date IS NOT NULL
            AND i.created_at >= :from_timestamp
            AND i.created_at < :to_exclusive
            AND i.created_at >= (
              i.source_observation_date::timestamp AT TIME ZONE 'Asia/Tashkent'
            )
        ),
        inspection_action_durations AS (
          SELECT extract(epoch FROM (a.created_at-i.completed_at))/3600.0 AS hours
          FROM corrective_actions a
          JOIN field_inspections i ON i.id=a.inspection_id
          JOIN scope_fields sf ON sf.field_id=a.field_id
          WHERE i.completed_at IS NOT NULL
            AND a.created_at >= :from_timestamp
            AND a.created_at < :to_exclusive
            AND a.created_at >= i.completed_at
        ),
        action_close_durations AS (
          SELECT extract(epoch FROM (a.closed_at-a.created_at))/3600.0 AS hours
          FROM corrective_actions a
          JOIN scope_fields sf ON sf.field_id=a.field_id
          WHERE a.closed_at IS NOT NULL
            AND a.closed_at >= :from_timestamp
            AND a.closed_at < :to_exclusive
            AND a.closed_at >= a.created_at
        ),
        duration_stats AS (
          SELECT
            (SELECT count(*) FROM attention_durations) AS attention_count,
            (SELECT percentile_cont(0.5) WITHIN GROUP (ORDER BY hours)
             FROM attention_durations) AS attention_median,
            (SELECT percentile_cont(0.9) WITHIN GROUP (ORDER BY hours)
             FROM attention_durations) AS attention_p90,
            (SELECT count(*) FROM inspection_action_durations)
              AS inspection_action_count,
            (SELECT percentile_cont(0.5) WITHIN GROUP (ORDER BY hours)
             FROM inspection_action_durations) AS inspection_action_median,
            (SELECT percentile_cont(0.9) WITHIN GROUP (ORDER BY hours)
             FROM inspection_action_durations) AS inspection_action_p90,
            (SELECT count(*) FROM action_close_durations) AS action_close_count,
            (SELECT percentile_cont(0.5) WITHIN GROUP (ORDER BY hours)
             FROM action_close_durations) AS action_close_median,
            (SELECT percentile_cont(0.9) WITHIN GROUP (ORDER BY hours)
             FROM action_close_durations) AS action_close_p90
        ),
        verification_outcomes AS (
          SELECT
            count(*) FILTER (WHERE v.result='improved') AS improved,
            count(*) FILTER (WHERE v.result='unchanged') AS unchanged,
            count(*) FILTER (WHERE v.result='worsened') AS worsened,
            count(*) FILTER (WHERE v.result='insufficient_data')
              AS insufficient_data
          FROM action_verification_requests v
          JOIN scope_fields sf ON sf.field_id=v.field_id
          WHERE v.status='resolved'
            AND v.resolved_at >= :from_timestamp
            AND v.resolved_at < :to_exclusive
        ),
        observations AS (
          SELECT 'ndvi'::text AS index_code, n.captured_date
          FROM ndvi_records n
          JOIN scope_fields sf ON sf.field_id=n.field_id
          WHERE n.captured_date <= :as_of_date
            AND {_ACCEPTED_NDVI_OBSERVATION}
          UNION ALL
          SELECT s.index_code, s.captured_date
          FROM satellite_index_records s
          JOIN scope_fields sf ON sf.field_id=s.field_id
          WHERE s.captured_date <= :as_of_date
            AND s.index_code IN ('savi','evi','ndmi','ndre')
            AND {_ACCEPTED_INDEX_OBSERVATION}
        ),
        latest_observations AS (
          SELECT index_code, max(captured_date) AS captured_date
          FROM observations
          GROUP BY index_code
        ),
        enterprise_inspections AS (
          SELECT sf.enterprise_id,
            count(*) FILTER (
              WHERE i.status IN {OPEN_INSPECTION_STATUSES}
            ) AS open_inspections,
            count(*) FILTER (
              WHERE i.status IN {OPEN_INSPECTION_STATUSES}
                AND i.assigned_to_id IS NULL
            ) AS unassigned_inspections,
            count(*) FILTER (
              WHERE i.status IN {OPEN_INSPECTION_STATUSES}
                AND i.due_date < :as_of_date
            ) AS overdue_inspections
          FROM scope_fields sf
          LEFT JOIN field_inspections i ON i.field_id=sf.field_id
          GROUP BY sf.enterprise_id
        ),
        enterprise_actions AS (
          SELECT sf.enterprise_id,
            count(lw.id) AS open_actions,
            count(lw.id) FILTER (WHERE lw.due_date < :as_of_date) AS overdue_actions
          FROM scope_fields sf
          LEFT JOIN live_work lw ON lw.field_id=sf.field_id
          GROUP BY sf.enterprise_id
        ),
        enterprise_awaiting AS (
          SELECT sf.enterprise_id,
            count(p.id) AS awaiting_verification
          FROM scope_fields sf
          LEFT JOIN agronomy_plans p ON p.field_id=sf.field_id
            AND p.status='pending_verification'
          GROUP BY sf.enterprise_id
        ),
        enterprise_verifications AS (
          SELECT sf.enterprise_id,
            count(*) FILTER (
              WHERE v.status='resolved' AND v.result='improved'
                AND v.resolved_at >= :from_timestamp
                AND v.resolved_at < :to_exclusive
            ) AS improved,
            count(*) FILTER (
              WHERE v.status='resolved' AND v.result='unchanged'
                AND v.resolved_at >= :from_timestamp
                AND v.resolved_at < :to_exclusive
            ) AS unchanged,
            count(*) FILTER (
              WHERE v.status='resolved' AND v.result='worsened'
                AND v.resolved_at >= :from_timestamp
                AND v.resolved_at < :to_exclusive
            ) AS worsened,
            count(*) FILTER (
              WHERE v.status='resolved' AND v.result='insufficient_data'
                AND v.resolved_at >= :from_timestamp
                AND v.resolved_at < :to_exclusive
            ) AS insufficient_data
          FROM scope_fields sf
          LEFT JOIN action_verification_requests v ON v.field_id=sf.field_id
          GROUP BY sf.enterprise_id
        ),
        enterprise_rows AS (
          SELECT sf.enterprise_id, min(sf.enterprise_name) AS enterprise_name,
            max(ei.open_inspections) AS open_inspections,
            max(ei.unassigned_inspections) AS unassigned_inspections,
            max(ei.overdue_inspections) AS overdue_inspections,
            max(ea.open_actions) AS open_actions,
            max(ea.overdue_actions) AS overdue_actions,
            max(aw.awaiting_verification) AS awaiting_verification,
            max(ev.improved) AS improved,
            max(ev.unchanged) AS unchanged,
            max(ev.worsened) AS worsened,
            max(ev.insufficient_data) AS insufficient_data
          FROM scope_fields sf
          LEFT JOIN enterprise_inspections ei
            ON ei.enterprise_id=sf.enterprise_id
          LEFT JOIN enterprise_actions ea
            ON ea.enterprise_id=sf.enterprise_id
          LEFT JOIN enterprise_verifications ev
            ON ev.enterprise_id=sf.enterprise_id
          LEFT JOIN enterprise_awaiting aw
            ON aw.enterprise_id=sf.enterprise_id
          GROUP BY sf.enterprise_id
        ),
        owner_rows AS (
          SELECT lw.assigned_to_id AS owner_id, u.full_name AS owner_name,
            count(*) AS unresolved_actions,
            count(*) FILTER (WHERE lw.due_date < :as_of_date) AS overdue_actions,
            min(lw.due_date) AS next_due_date
          FROM live_work lw
          JOIN users u ON u.id=lw.assigned_to_id
          GROUP BY lw.assigned_to_id,u.full_name
        )
        SELECT ins.open_inspections, ins.unassigned_inspections,
          ins.overdue_inspections, act.open_actions, act.overdue_actions,
          ver.awaiting_verification,
          ds.attention_count, ds.attention_median, ds.attention_p90,
          ds.inspection_action_count, ds.inspection_action_median,
          ds.inspection_action_p90, ds.action_close_count,
          ds.action_close_median, ds.action_close_p90,
          outcomes.improved, outcomes.unchanged, outcomes.worsened,
          outcomes.insufficient_data,
          COALESCE((
            SELECT json_object_agg(index_code,captured_date)
            FROM latest_observations
          ), '{{}}'::json) AS latest_observations,
          COALESCE((
            SELECT json_agg(json_build_object(
              'enterprise_id',enterprise_id,
              'enterprise_name',enterprise_name,
              'open_inspections',open_inspections,
              'unassigned_inspections',unassigned_inspections,
              'overdue_inspections',overdue_inspections,
              'open_actions',open_actions,
              'overdue_actions',overdue_actions,
              'awaiting_verification',awaiting_verification,
              'verification_outcomes',json_build_object(
                'improved',improved,
                'unchanged',unchanged,
                'worsened',worsened,
                'insufficient_data',insufficient_data
              )
            ) ORDER BY enterprise_name,enterprise_id)
            FROM enterprise_rows
          ), '[]'::json) AS enterprises,
          COALESCE((
            SELECT json_agg(json_build_object(
              'owner_id',owner_id,
              'owner_name',owner_name,
              'unresolved_actions',unresolved_actions,
              'overdue_actions',overdue_actions,
              'next_due_date',next_due_date
            ) ORDER BY overdue_actions DESC,unresolved_actions DESC,
              owner_name,owner_id)
            FROM owner_rows
          ), '[]'::json) AS owners
        FROM inspection_snapshot ins
        CROSS JOIN action_snapshot act
        CROSS JOIN verification_snapshot ver
        CROSS JOIN duration_stats ds
        CROSS JOIN verification_outcomes outcomes
        """
    )


def _attention_metrics(queue):
    items = queue["items"]
    by_enterprise = {}
    for item in items:
        enterprise_id = item["field"]["enterprise_id"]
        current = by_enterprise.setdefault(
            enterprise_id,
            {
                "attention_fields_now": 0,
                "attention_critical": 0,
                "attention_high": 0,
                "attention_medium": 0,
            },
        )
        current["attention_fields_now"] += 1
        current[f"attention_{item['priority']}"] += 1
    return {
        "summary": queue["summary"],
        "stale_fields": sum(
            item["spectral_summary"]["data_status"] == "stale"
            for item in items
        ),
        "no_data_fields": sum(
            item["spectral_summary"]["data_status"] == "no_data"
            for item in items
        ),
        "low_confidence_fields": sum(
            item["spectral_summary"]["overall_confidence"]
            in {"low", "insufficient"}
            for item in items
        ),
        "by_enterprise": by_enterprise,
    }


def overview(
    db,
    user,
    *,
    date_from=None,
    date_to=None,
    enterprise_id=None,
    attention_lookback_days=180,
):
    scope = resolve_scope(user, enterprise_id)
    window = resolve_window(date_from, date_to)
    try:
        queue = build_attention_queue(
            db,
            user,
            enterprise_id=scope.enterprise_id,
            date_to=window.date_to,
            lookback_days=attention_lookback_days,
            min_priority="medium",
            limit=MAX_SCOPE_FIELDS,
        )
        attention = _attention_metrics(queue)
        tenant_clause = (
            "WHERE f.enterprise_id=:enterprise_id"
            if scope.enterprise_id is not None
            else ""
        )
        params = {
            "enterprise_id": scope.enterprise_id,
            "as_of_date": window.date_to,
            "from_timestamp": window.from_timestamp,
            "to_exclusive": window.to_exclusive,
        }
        row = _one(db.execute(_overview_sql(tenant_clause), params))
        if row is None:
            raise HTTPException(500, "Executive aggregate query returned no row")
        enterprises = [dict(item) for item in (row["enterprises"] or [])]
        for item in enterprises:
            item.update(
                attention["by_enterprise"].get(
                    item["enterprise_id"],
                    {
                        "attention_fields_now": 0,
                        "attention_critical": 0,
                        "attention_high": 0,
                        "attention_medium": 0,
                    },
                )
            )
        latest = {
            code: None
            for code in ("ndvi", "savi", "evi", "ndmi", "ndre")
        }
        latest.update(row["latest_observations"] or {})
        return {
            "definitions_version": DEFINITIONS_VERSION,
            "generated_at": datetime.now(TASHKENT),
            "timezone": "Asia/Tashkent",
            "scope": {
                "role": scope.role,
                "enterprise_id": scope.enterprise_id,
            },
            "date_range": {
                "from": window.date_from,
                "to": window.date_to,
                "inclusive": True,
            },
            "backlog": {
                "attention_fields_now": _integer(
                    attention["summary"]["attention_fields"]
                ),
                "attention_critical": _integer(
                    attention["summary"]["critical"]
                ),
                "attention_high": _integer(attention["summary"]["high"]),
                "attention_medium": _integer(attention["summary"]["medium"]),
                "open_inspections": _integer(row["open_inspections"]),
                "unassigned_inspections": _integer(
                    row["unassigned_inspections"]
                ),
                "overdue_inspections": _integer(row["overdue_inspections"]),
                "open_actions": _integer(row["open_actions"]),
                "overdue_actions": _integer(row["overdue_actions"]),
                "awaiting_verification": _integer(
                    row["awaiting_verification"]
                ),
            },
            "cycle_times": {
                "attention_signal_to_inspection_hours": _duration(
                    row,
                    "attention",
                ),
                "inspection_to_action_hours": _duration(
                    row,
                    "inspection_action",
                ),
                "action_to_close_hours": _duration(row, "action_close"),
            },
            "verification_outcomes": {
                key: _integer(row[key])
                for key in (
                    "improved",
                    "unchanged",
                    "worsened",
                    "insufficient_data",
                )
            },
            "data_quality": {
                "attention_stale_fields": attention["stale_fields"],
                "attention_no_data_fields": attention["no_data_fields"],
                "attention_low_confidence_fields": attention[
                    "low_confidence_fields"
                ],
                "latest_observation_by_index": latest,
            },
            "enterprises": enterprises,
            "owners": [dict(item) for item in (row["owners"] or [])],
            "limitations": list(LIMITATIONS),
        }
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise


def _accountability_query(
    kind,
    *,
    tenant_scoped,
    owner_filtered,
    count_only,
):
    tenant = " AND f.enterprise_id=:enterprise_id" if tenant_scoped else ""
    pagination = "" if count_only else " LIMIT :limit OFFSET :offset"
    if kind in {"unassigned_inspections", "overdue_inspections"}:
        condition = (
            f"i.status IN {OPEN_INSPECTION_STATUSES} AND i.assigned_to_id IS NULL"
            if kind == "unassigned_inspections"
            else (
                f"i.status IN {OPEN_INSPECTION_STATUSES} "
                "AND i.due_date < :as_of_date"
            )
        )
        if count_only:
            return text(
                "SELECT count(*) AS total FROM field_inspections i "
                "JOIN fields f ON f.id=i.field_id "
                f"WHERE {condition}{tenant}"
            )
        return text(
            "SELECT i.id,i.id AS inspection_id,NULL::integer AS action_id,"
            "NULL::bigint AS plan_id,"
            "i.field_id,f.name AS field_name,f.enterprise_id,"
            "e.name AS enterprise_name,i.assigned_to_id AS owner_id,"
            "u.full_name AS owner_name,i.title AS description,i.due_date,"
            "i.status,i.version,NULL::integer AS verification_id,"
            "NULL::text AS verification_status "
            "FROM field_inspections i JOIN fields f ON f.id=i.field_id "
            "JOIN enterprises e ON e.id=f.enterprise_id "
            "LEFT JOIN users u ON u.id=i.assigned_to_id "
            f"WHERE {condition}{tenant} "
            "ORDER BY i.due_date ASC NULLS LAST,i.created_at ASC,i.id ASC"
            + pagination
        )
    if kind in {"open_actions", "overdue_actions"}:
        condition = f"{LIVE_PLAN} AND {ACTIVE_WORK}"
        if kind == "overdue_actions":
            condition += f" AND {WORK_DUE_DATE} < :as_of_date"
        owner = " AND w.assigned_to_id=:owner_id" if owner_filtered else ""
        source = (
            "FROM agronomy_work_items w "
            "JOIN agronomy_plans p ON p.id=w.plan_id AND p.enterprise_id=w.enterprise_id "
            "AND p.cycle=w.cycle "
            "JOIN fields f ON f.id=w.field_id AND f.enterprise_id=w.enterprise_id "
        )
        if count_only:
            return text(f"SELECT count(*) AS total {source}WHERE {condition}{tenant}{owner}")
        return text(
            "SELECT w.id,w.inspection_id,w.id AS action_id,w.plan_id,w.field_id,"
            "f.name AS field_name,f.enterprise_id,e.name AS enterprise_name,"
            "w.assigned_to_id AS owner_id,u.full_name AS owner_name,"
            f"w.instruction AS description,{WORK_DUE_DATE} AS due_date,"
            "w.status,w.version,NULL::bigint AS verification_id,"
            "p.verification_status "
            + source
            + "JOIN enterprises e ON e.id=f.enterprise_id "
            "LEFT JOIN users u ON u.id=w.assigned_to_id "
            f"WHERE {condition}{tenant}{owner} "
            "ORDER BY w.due_at ASC NULLS LAST,w.created_at ASC,w.id ASC"
            + pagination
        )
    owner = (
        " AND EXISTS (SELECT 1 FROM agronomy_work_items ow WHERE ow.plan_id=p.id "
        "AND ow.cycle=p.cycle AND ow.assigned_to_id=:owner_id)"
        if owner_filtered else ""
    )
    source = (
        "FROM agronomy_plans p "
        "JOIN fields f ON f.id=p.field_id AND f.enterprise_id=p.enterprise_id "
    )
    condition = "p.status='pending_verification'"
    if count_only:
        return text(f"SELECT count(*) AS total {source}WHERE {condition}{tenant}{owner}")
    return text(
        "SELECT p.id,p.inspection_id,NULL::bigint AS action_id,p.id AS plan_id,p.field_id,"
        "f.name AS field_name,f.enterprise_id,e.name AS enterprise_name,"
        "p.approved_by_id AS owner_id,u.full_name AS owner_name,"
        "p.decision AS description,NULL::date AS due_date,p.status,p.version,"
        "(SELECT v.id FROM agronomy_verifications v WHERE v.plan_id=p.id "
        "ORDER BY v.id DESC LIMIT 1) AS verification_id,"
        "p.verification_status "
        + source
        + "JOIN enterprises e ON e.id=f.enterprise_id "
        "LEFT JOIN users u ON u.id=p.approved_by_id "
        f"WHERE {condition}{tenant}{owner} "
        "ORDER BY p.completed_at ASC NULLS LAST,p.id ASC"
        + pagination
    )


def accountability(
    db,
    user,
    *,
    kind,
    date_from=None,
    date_to=None,
    enterprise_id=None,
    owner_id=None,
    limit=50,
    offset=0,
):
    if kind not in ACCOUNTABILITY_KINDS:
        raise HTTPException(422, "Unknown accountability kind")
    if owner_id is not None and kind in {
        "unassigned_inspections",
        "overdue_inspections",
    }:
        raise HTTPException(422, "owner_id applies only to action queues")
    scope = resolve_scope(user, enterprise_id)
    window = resolve_window(date_from, date_to)
    params = {
        "enterprise_id": scope.enterprise_id,
        "owner_id": owner_id,
        "as_of_date": window.date_to,
        "limit": limit,
        "offset": offset,
    }
    tenant_scoped = scope.enterprise_id is not None
    owner_filtered = owner_id is not None
    try:
        total_row = _one(
            db.execute(
                _accountability_query(
                    kind,
                    tenant_scoped=tenant_scoped,
                    owner_filtered=owner_filtered,
                    count_only=True,
                ),
                params,
            )
        )
        rows = _all(
            db.execute(
                _accountability_query(
                    kind,
                    tenant_scoped=tenant_scoped,
                    owner_filtered=owner_filtered,
                    count_only=False,
                ),
                params,
            )
        )
        return {
            "definitions_version": DEFINITIONS_VERSION,
            "generated_at": datetime.now(TASHKENT),
            "timezone": "Asia/Tashkent",
            "scope": {
                "role": scope.role,
                "enterprise_id": scope.enterprise_id,
            },
            "date_range": {
                "from": window.date_from,
                "to": window.date_to,
                "inclusive": True,
            },
            "kind": kind,
            "owner_id": owner_id,
            "total": _integer(total_row["total"]),
            "limit": limit,
            "offset": offset,
            "items": [dict(row) for row in rows],
        }
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise


def executive_workbook(payload):
    """Build one deterministic XLSX from the exact overview payload."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    rows = [
        ("Definitions version", payload["definitions_version"]),
        ("Generated at", payload["generated_at"].isoformat()),
        ("Timezone", payload["timezone"]),
        ("Date from", payload["date_range"]["from"].isoformat()),
        ("Date to", payload["date_range"]["to"].isoformat()),
    ]
    rows.extend(
        (name, value)
        for name, value in payload["backlog"].items()
    )
    for row in rows:
        summary.append(row)

    enterprises = workbook.create_sheet("Enterprises")
    enterprise_headers = [
        "enterprise_id",
        "enterprise_name",
        "attention_fields_now",
        "open_inspections",
        "unassigned_inspections",
        "overdue_inspections",
        "open_actions",
        "overdue_actions",
        "awaiting_verification",
    ]
    enterprises.append(enterprise_headers)
    for item in payload["enterprises"]:
        enterprises.append([item.get(key) for key in enterprise_headers])

    owners = workbook.create_sheet("Owners")
    owner_headers = [
        "owner_id",
        "owner_name",
        "unresolved_actions",
        "overdue_actions",
        "next_due_date",
    ]
    owners.append(owner_headers)
    for item in payload["owners"]:
        owners.append([item.get(key) for key in owner_headers])

    definitions = workbook.create_sheet("Definitions")
    definitions.append(["Metric", "Value"])
    for name, metric in payload["cycle_times"].items():
        definitions.append(
            [
                name,
                (
                    f"n={metric['sample_count']}; "
                    f"median={metric['median_hours']}; "
                    f"p90={metric['p90_hours']}"
                ),
            ]
        )
    for limitation in payload["limitations"]:
        definitions.append(["limitation", limitation])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
