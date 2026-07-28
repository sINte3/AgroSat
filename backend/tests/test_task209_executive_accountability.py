from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from main import app
from schemas.executive_accountability import (
    AccountabilityResponse,
    ExecutiveOverviewResponse,
)
from services import executive_accountability as service


class MappingResult:
    def __init__(self, *, one=None, all_rows=None):
        self.one = one
        self.all_rows = list(all_rows or [])

    def mappings(self):
        return self

    def first(self):
        return self.one

    def all(self):
        return self.all_rows


class RecordingDB:
    def __init__(self, results):
        self.results = list(results)
        self.calls = []
        self.rollbacks = 0

    def execute(self, statement, params):
        self.calls.append((str(statement), dict(params)))
        return self.results.pop(0)

    def rollback(self):
        self.rollbacks += 1


def actor(role="manager", enterprise_id=5, user_id=7):
    return SimpleNamespace(
        id=user_id,
        role=role,
        enterprise_id=enterprise_id,
        is_active=True,
    )


def attention_queue():
    return {
        "summary": {
            "fields_evaluated": 2,
            "attention_fields": 2,
            "critical": 1,
            "high": 0,
            "medium": 1,
            "low": 0,
            "returned": 2,
        },
        "items": [
            {
                "priority": "critical",
                "field": {"enterprise_id": 5},
                "spectral_summary": {
                    "data_status": "stale",
                    "overall_confidence": "low",
                },
            },
            {
                "priority": "medium",
                "field": {"enterprise_id": 5},
                "spectral_summary": {
                    "data_status": "no_data",
                    "overall_confidence": "insufficient",
                },
            },
        ],
    }


def overview_row():
    return {
        "open_inspections": 4,
        "unassigned_inspections": 2,
        "overdue_inspections": 1,
        "open_actions": 3,
        "overdue_actions": 2,
        "awaiting_verification": 1,
        "attention_count": 3,
        "attention_median": 12.25,
        "attention_p90": 24.5,
        "inspection_action_count": 2,
        "inspection_action_median": 4,
        "inspection_action_p90": 7.5,
        "action_close_count": 1,
        "action_close_median": 30,
        "action_close_p90": 30,
        "improved": 2,
        "unchanged": 1,
        "worsened": 1,
        "insufficient_data": 3,
        "latest_observations": {
            "ndvi": date(2026, 7, 20),
            "ndmi": date(2026, 7, 19),
        },
        "enterprises": [
            {
                "enterprise_id": 5,
                "enterprise_name": "Synthetic Enterprise",
                "open_inspections": 4,
                "unassigned_inspections": 2,
                "overdue_inspections": 1,
                "open_actions": 3,
                "overdue_actions": 2,
                "awaiting_verification": 1,
                "verification_outcomes": {
                    "improved": 2,
                    "unchanged": 1,
                    "worsened": 1,
                    "insufficient_data": 3,
                },
            }
        ],
        "owners": [
            {
                "owner_id": 8,
                "owner_name": "Synthetic Agronomist",
                "unresolved_actions": 3,
                "overdue_actions": 2,
                "next_due_date": date(2026, 7, 30),
            }
        ],
    }


def build_overview(db=None, user=None):
    db = db or RecordingDB([MappingResult(one=overview_row())])
    with patch.object(
        service,
        "build_attention_queue",
        return_value=attention_queue(),
    ) as attention:
        value = service.overview(
            db,
            user or actor(),
            date_from=date(2026, 7, 1),
            date_to=date(2026, 7, 28),
            attention_lookback_days=120,
        )
    return value, db, attention


def test_management_scope_is_server_enforced():
    assert service.resolve_scope(actor("admin", None), None).enterprise_id is None
    assert service.resolve_scope(actor("admin", None), 9).enterprise_id == 9
    assert service.resolve_scope(actor(), None).enterprise_id == 5
    assert service.resolve_scope(actor(), 5).enterprise_id == 5
    for denied in (
        actor("agronomist"),
        actor("viewer"),
        actor("unknown"),
        actor("manager", None),
    ):
        with pytest.raises(HTTPException) as caught:
            service.resolve_scope(denied)
        assert caught.value.status_code == 403
    with pytest.raises(HTTPException) as caught:
        service.resolve_scope(actor(), 6)
    assert caught.value.status_code == 403


def test_date_window_is_inclusive_tashkent_and_bounded():
    window = service.resolve_window(
        date(2026, 1, 1),
        date(2026, 12, 31),
    )
    assert window.from_timestamp.utcoffset().total_seconds() == 5 * 3600
    assert window.to_exclusive.date() == date(2027, 1, 1)
    with pytest.raises(HTTPException):
        service.resolve_window(date(2026, 2, 1), date(2026, 1, 1))
    with pytest.raises(HTTPException):
        service.resolve_window(date(2025, 1, 1), date(2026, 1, 2))


def test_overview_has_one_operational_query_and_reuses_attention_contract():
    payload, db, attention = build_overview()
    assert len(db.calls) == 1
    sql, params = db.calls[0]
    assert "WHERE f.enterprise_id=:enterprise_id" in sql
    assert params["enterprise_id"] == 5
    assert "percentile_cont(0.5)" in sql
    assert "percentile_cont(0.9)" in sql
    assert "action_verification_requests" in sql
    assert "source_observation_date::timestamp AT TIME ZONE 'Asia/Tashkent'" in sql
    attention.assert_called_once()
    assert attention.call_args.kwargs["limit"] == service.MAX_SCOPE_FIELDS
    assert attention.call_args.kwargs["enterprise_id"] == 5
    assert payload["backlog"]["attention_fields_now"] == 2
    assert payload["data_quality"]["attention_stale_fields"] == 1
    assert payload["data_quality"]["attention_no_data_fields"] == 1
    assert payload["data_quality"]["attention_low_confidence_fields"] == 2
    assert payload["enterprises"][0]["attention_fields_now"] == 2
    assert payload["cycle_times"]["attention_signal_to_inspection_hours"] == {
        "sample_count": 3,
        "median_hours": 12.25,
        "p90_hours": 24.5,
    }
    ExecutiveOverviewResponse.model_validate(payload)


def test_admin_global_query_has_no_tenant_predicate():
    db = RecordingDB([MappingResult(one=overview_row())])
    build_overview(db, actor("admin", None))
    assert "WHERE f.enterprise_id=:enterprise_id" not in db.calls[0][0]


def accountability_row():
    return {
        "id": 401,
        "inspection_id": 101,
        "action_id": 401,
        "field_id": 11,
        "field_name": "Synthetic Field",
        "enterprise_id": 5,
        "enterprise_name": "Synthetic Enterprise",
        "owner_id": 8,
        "owner_name": "Synthetic Agronomist",
        "description": "Bounded corrective action",
        "due_date": date(2026, 7, 20),
        "status": "open",
        "version": 2,
        "verification_id": None,
        "verification_status": None,
    }


@pytest.mark.parametrize(
    "kind,required",
    [
        ("unassigned_inspections", "i.assigned_to_id IS NULL"),
        ("overdue_inspections", "i.due_date < :as_of_date"),
        ("overdue_actions", "a.due_date < :as_of_date"),
        ("awaiting_verification", "v.status='awaiting_observation'"),
    ],
)
def test_accountability_is_exactly_two_bounded_queries(kind, required):
    row = accountability_row()
    if kind in {"unassigned_inspections", "overdue_inspections"}:
        row = {
            **row,
            "id": 101,
            "action_id": None,
            "owner_id": None,
            "owner_name": None,
            "description": "Bounded inspection",
        }
    db = RecordingDB(
        [
            MappingResult(one={"total": 1}),
            MappingResult(all_rows=[row]),
        ]
    )
    payload = service.accountability(
        db,
        actor(),
        kind=kind,
        date_to=date(2026, 7, 28),
        limit=25,
        offset=50,
    )
    assert len(db.calls) == 2
    assert all(required in sql for sql, _ in db.calls)
    assert all("f.enterprise_id=:enterprise_id" in sql for sql, _ in db.calls)
    assert "LIMIT :limit OFFSET :offset" in db.calls[1][0]
    assert db.calls[1][1]["limit"] == 25
    assert db.calls[1][1]["offset"] == 50
    assert payload["total"] == 1
    AccountabilityResponse.model_validate(payload)


def test_owner_filter_is_action_only_and_inside_both_queries():
    with pytest.raises(HTTPException) as caught:
        service.accountability(
            RecordingDB([]),
            actor(),
            kind="unassigned_inspections",
            owner_id=8,
        )
    assert caught.value.status_code == 422
    db = RecordingDB(
        [
            MappingResult(one={"total": 0}),
            MappingResult(all_rows=[]),
        ]
    )
    service.accountability(
        db,
        actor(),
        kind="overdue_actions",
        owner_id=8,
    )
    assert all("a.owner_id=:owner_id" in sql for sql, _ in db.calls)


def test_workbook_reconciles_with_overview_totals():
    payload, _, _ = build_overview()
    content = service.executive_workbook(payload)
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == [
        "Summary",
        "Enterprises",
        "Owners",
        "Definitions",
    ]
    summary = dict(workbook["Summary"].iter_rows(values_only=True))
    assert summary["attention_fields_now"] == payload["backlog"][
        "attention_fields_now"
    ]
    assert summary["overdue_actions"] == payload["backlog"][
        "overdue_actions"
    ]
    enterprise_rows = list(
        workbook["Enterprises"].iter_rows(values_only=True)
    )
    assert enterprise_rows[1][0] == 5
    assert enterprise_rows[1][2] == 2


def test_openapi_contract_has_three_authenticated_management_routes():
    paths = app.openapi()["paths"]
    for path in (
        "/api/executive/overview",
        "/api/executive/accountability",
        "/api/executive/export.xlsx",
    ):
        operation = paths[path]["get"]
        assert operation["security"] == [{"OAuth2PasswordBearer": []}]


def test_service_has_no_orm_relationship_or_unbounded_page_contract():
    source = (
        __import__("pathlib").Path(service.__file__).read_text(encoding="utf-8")
    )
    assert "relationship(" not in source
    assert ".query(" not in source
    assert " LIMIT :limit OFFSET :offset" in source
    assert "MAX_SCOPE_FIELDS" in source
